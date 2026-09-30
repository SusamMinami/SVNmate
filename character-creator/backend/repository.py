"""Lossless positional CSV reading and read-only xlsm inspection."""
import csv
from collections import Counter
import hashlib
import io
import json
import os
from pathlib import Path

from .catalog import TABLES, EDITABLE, SKILL_REFS
from .authoring_schema import ARRAYS


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def member(value):
    return text(value).lstrip("#&").strip()


def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def suggested_doc():
    path = Path(os.environ.get("APPDATA", "")) / "Shot Sandbox" / "desktop-state.json"
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        # Only inspect directory settings, never return unrelated credentials.
        def visit(value):
            if isinstance(value, dict):
                for key, item in value.items():
                    if isinstance(item, str) and "doc" in key.lower():
                        candidate = Path(item)
                        if (candidate / "csvdir").is_dir():
                            return str(candidate)
                    found = visit(item)
                    if found:
                        return found
            return ""
        return visit(data)
    except (OSError, ValueError):
        return ""


class Table:
    def __init__(self, path, prefix):
        raw = Path(path).read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError(f"无法识别 CSV 编码：{path.name}")
        rows = list(csv.reader(io.StringIO(decoded)))
        if len(rows) < 2:
            raise ValueError(f"缺少双表头：{path.name}")
        self.headers = [member(h) for h in rows[0]]
        self.labels = rows[1]
        self.path = Path(path)
        self.prefix = prefix
        self.rows = [r + [""] * max(0, len(self.headers) - len(r)) for r in rows[2:]
                     if r and r[0].strip() and not r[0].startswith("#")]
        self.index = {}
        for row in self.rows:
            self.index.setdefault(row[0], []).append(row)

    def row(self, key):
        matches = self.index.get(str(key), [])
        if len(matches) != 1:
            raise ValueError(f"{self.path.stem} ID {key} 找到 {len(matches)} 行，要求唯一")
        return matches[0]

    def column(self, field):
        key = field if "." in field else f"{self.prefix}.{field}"
        hits = [i for i, h in enumerate(self.headers) if h == key]
        if len(hits) != 1:
            raise ValueError(f"成员 {key} 不唯一或不存在")
        return hits[0]

    def get(self, row, field):
        try:
            return row[self.column(field)]
        except ValueError:
            return ""

    def positions(self, field):
        start = self.column(field)
        if self.headers[start] in ARRAYS:
            # Array terminators can have blank or repeated nested member names.
            for end in range(start + 1, len(self.headers)):
                if self.rows and all(r[end] == "}" for r in self.rows[:20]):
                    return list(range(start, end + 1))
            raise ValueError(f"{field} 缺少一致的数组结束列")
        return [start]

    def fields(self, row, table_key):
        fields = []
        for i, header in enumerate(self.headers):
            name = header.split(".", 1)[-1]
            if name not in EDITABLE.get(table_key, set()):
                continue
            positions = self.positions(name)
            values = [row[p] for p in positions]
            fields.append({"name": name, "label": self.labels[i].lstrip("#") or name,
                           "values": values, "labels": [self.labels[p] for p in positions],
                           "member": header})
        return fields


class Repository:
    def __init__(self, doc):
        self.doc = Path(doc).resolve()
        if not (self.doc / "csvdir").is_dir() or not (self.doc / "xlsdir").is_dir():
            raise ValueError("请选择同时包含 csvdir 和 xlsdir 的 doc 目录")
        self.tables = {}
        self.missing = []
        for key, (stem, prefix, _) in TABLES.items():
            path = self.doc / "csvdir" / f"{stem}.csv"
            if path.exists():
                self.tables[key] = Table(path, prefix)
            else:
                self.missing.append(stem)
        if "career" not in self.tables:
            raise ValueError("缺少 z职业配置表.csv")
        self.snapshot = hashlib.sha256("".join(
            k + fingerprint(t.path) for k, t in sorted(self.tables.items())
        ).encode()).hexdigest()[:16]

    def workbook(self, key):
        stem = TABLES[key][0]
        matches = [p for p in (self.doc / "xlsdir").rglob(f"{stem}.xls*")
                   if p.suffix.lower() in {".xlsx", ".xlsm"} and not p.name.startswith("~$")]
        if len(matches) != 1:
            raise ValueError(f"{stem} 找到 {len(matches)} 个工作簿，须保持唯一来源")
        return matches[0].resolve()

    def list_careers(self):
        table = self.tables["career"]
        create = self.tables.get("create")
        base_ids = {create.get(r, "careerid") for r in create.rows} if create else set()
        return [{"id": row[0], "name": table.get(row, "name"),
                 "family": table.get(row, "bp").split("/")[0],
                 "role": ("新手本职业" if table.get(row, "IsIntroChar") == "1" else
                          "创角基础职业" if row[0] in base_ids else "职业 / 转职分支"),
                 "enabled": table.get(row, "onoffswitch") == "1"}
                for row in table.rows]

    def detail(self, career_id):
        primary = self.tables["career"]
        career = primary.row(career_id)
        sections = []
        for key in EDITABLE:
            table = self.tables.get(key)
            if table and str(career_id) in table.index:
                row = table.row(career_id)
                sections.append({"key": key, "title": TABLES[key][2],
                                 "source": table.path.name, "fields": table.fields(row, key)})
        skill_ids = set()
        for p in primary.positions("initial_skill"):
            value = career[p]
            if ";" in value:
                skill_ids.add(value.split(";")[0])
        for h, value in zip(primary.headers, career):
            candidate = value.split(";")[0]
            if h.split(".")[-1] in SKILL_REFS and candidate.isdigit() and int(candidate) > 0:
                skill_ids.add(candidate)
        skills = []
        table = self.tables.get("skill")
        tree = self.tables.get("tree")
        if table:
            for row in table.rows:
                explicit = row[0] in skill_ids
                if explicit or table.get(row, "career") == str(career_id):
                    links = [] if not tree else [
                        {"id": r[0], "name": tree.get(r, "skillname"),
                         "maxLevel": tree.get(r, "skillmaxlv")}
                        for r in tree.rows if tree.get(r, "skillid") == row[0]]
                    skills.append({"id": row[0], "name": table.get(row, "skillname"),
                                   "ability": table.get(row, "skillgameplay"),
                                   "cd": table.get(row, "skillcd"), "tree": links,
                                   "basis": "职业显式引用" if explicit else "Skill.career"})
        create = self.tables.get("create")
        creation = [] if not create else [
            {h: r[i] for i, h in enumerate(create.headers) if h}
            for r in create.rows if create.get(r, "careerid") == str(career_id)]
        return {"id": str(career_id), "name": primary.get(career, "name"),
                "family": primary.get(career, "bp").split("/")[0],
                "sections": sections, "skills": skills, "creation": creation,
                "missingSkills": sorted(skill_ids - set(table.index if table else {}))}


def inspect_workbook(path, table, source_id, target_id, edits, cache=None):
    """Map by unique full member name, not CSV offsets. Capture row baseline."""
    import openpyxl
    book = None
    try:
        id_member = f"{table.prefix}.id"
        cache_key = (str(path), id_member)
        cached = cache.get(cache_key) if cache is not None else None
        if cached is None:
            before = fingerprint(path)
            book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
            matches = []
            expected = Counter(h for h in table.headers if "." in h)
            for sheet in book.worksheets:
                head = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                headers = [member(v) for v in head]
                available = Counter(headers)
                if id_member in headers and all(available[h] >= count for h, count in expected.items()):
                    matches.append((sheet, headers))
            if len(matches) != 1:
                raise ValueError(f"{path.name}：无法唯一定位完整字段匹配的 {id_member} 工作表")
            sheet, headers = matches[0]
            rows = list(sheet.iter_rows(min_row=3, values_only=True))
            cached = (sheet.title, headers, rows, before)
            if fingerprint(path) != before:
                raise ValueError("工作簿在读取期间已改变，请重新审核")
            if cache is not None:
                cache[cache_key] = cached
        sheet_title, headers, rows, book_hash = cached
        if headers.count(id_member) != 1:
            raise ValueError("ID 列不唯一")
        id_col = headers.index(id_member)
        sources, targets = [], []
        # One streaming pass; individual read-only cell lookups are expensive.
        for number, values in enumerate(rows, 3):
            if text(values[id_col]) == str(source_id):
                sources.append((number, list(values)))
            if text(values[id_col]) == str(target_id):
                targets.append(number)
        if len(sources) != 1:
            raise ValueError(f"{path.name}：ID {source_id} 多版本或不存在，暂不支持自动选择")
        if str(source_id) != str(target_id) and targets:
            raise ValueError(f"{path.name}：目标 ID {target_id} 已存在")
        row, values = sources[0]
        changes = []
        csv_row = table.row(source_id)
        for name, new_values in edits.items():
            full = name if "." in name else f"{table.prefix}.{name}"
            if headers.count(full) != 1:
                raise ValueError(f"{path.name}：成员 {full} 不唯一")
            column = headers.index(full)
            positions = table.positions(name)
            if len(positions) > 1:
                if headers[column:column + len(positions)] != [table.headers[p] for p in positions]:
                    raise ValueError(f"{full}：Excel 与 CSV 数组结构不同")
            for offset, value in enumerate(new_values):
                old = text(values[column + offset])
                if old != csv_row[positions[offset]]:
                    raise ValueError(f"{full}：CSV 与 Excel 原值不同，请先核对导表快照")
                if old.startswith("="):
                    raise ValueError(f"{full}：公式单元格不能由基础面板覆盖")
                if old != value:
                    changes.append({"column": column + offset + 1, "member": full,
                                    "slot": offset, "before": old, "after": value})
        return {"path": str(path), "sheet": sheet_title, "row": row,
                "idColumn": id_col + 1, "sourceId": str(source_id), "targetId": str(target_id),
                "headers": headers, "baseline": [text(v) for v in values],
                "changes": changes, "fingerprint": book_hash,
                "clone": str(source_id) != str(target_id)}
    finally:
        if book:
            book.close()
