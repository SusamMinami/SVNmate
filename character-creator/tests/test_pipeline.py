import csv
from pathlib import Path
import tempfile
import unittest

import openpyxl

from backend.repository import Table, fingerprint
from backend.service import Service


def fixture(root):
    csvdir = root / "csvdir"
    xlsdir = root / "xlsdir"
    csvdir.mkdir()
    xlsdir.mkdir()
    headers = ["##&CareerInfor.id", "CareerInfor.name", "CareerInfor.bp",
               "CareerInfor.initial_skill", "", "", "", "CareerInfor.onoffswitch"]
    labels = ["##职业ID", "职业名", "蓝图", "初始技能", "槽位1", "槽位2", "", "开放"]
    row = ["100", "合成职业", "Synthetic/BP_Test", "{", "100001;1", "", "}", "1"]
    with (csvdir / "z职业配置表.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows([headers, labels, row])
    with (csvdir / "j技能表.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows([["##&Skill.id", "Skill.skillname", "Skill.career"],
                                      ["##id", "名称", "职业"], ["100001", "合成技能", "-1"]])
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["版本"] + [h.lstrip("#&") for h in headers] + ["公式保留"])
    sheet.append(["版本"] + labels + ["公式保留"])
    sheet.append(["测试"] + row + ["=1+1"])
    path = xlsdir / "z职业配置表.xlsx"
    book.save(path)
    book.close()
    return path


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.path = fixture(self.root)
        self.calls = []
        self.service = Service(writer=lambda plans, mode: self.calls.append((plans, mode)) or {"status": "written"})
        self.service.load(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def request(self, **kwargs):
        return {"sourceId": "100", "edits": {"career": {"name": ["测试改名"]}}, **kwargs}

    def change_workbook(self, operation):
        book = openpyxl.load_workbook(self.path)
        operation(book.active)
        book.save(self.path)
        book.close()

    def test_array_empty_slots_and_explicit_cross_career_skill(self):
        detail = self.service.detail("100")
        field = next(f for f in detail["sections"][0]["fields"] if f["name"] == "initial_skill")
        self.assertEqual(field["values"], ["{", "100001;1", "", "}"])
        self.assertEqual(detail["skills"][0]["basis"], "职业显式引用")

    def test_member_mapping_uses_excel_version_offset_and_preserves_formula(self):
        before = fingerprint(self.path)
        plan = self.service.prepare(self.request())["plans"][0]
        self.assertEqual(plan["changes"][0]["column"], 3)
        self.assertEqual(plan["changes"][0]["before"], "合成职业")
        self.assertEqual(fingerprint(self.path), before)

    def test_csv_workbook_divergence_blocks(self):
        self.change_workbook(lambda sheet: setattr(sheet.cell(3, 3), "value", "不同基线"))
        with self.assertRaisesRegex(ValueError, "原值不同"):
            self.service.prepare(self.request())

    def test_duplicate_version_blocks(self):
        self.change_workbook(lambda sheet: sheet.append([c.value for c in sheet[3]]))
        with self.assertRaisesRegex(ValueError, "多版本"):
            self.service.prepare(self.request())

    def test_duplicate_workbook_blocks(self):
        (self.root / "xlsdir" / "duplicate").mkdir()
        (self.root / "xlsdir" / "duplicate" / self.path.name).write_bytes(self.path.read_bytes())
        with self.assertRaisesRegex(ValueError, "2 个"):
            self.service.prepare(self.request())

    def test_duplicate_member_blocks(self):
        self.change_workbook(lambda sheet: setattr(sheet.cell(1, 10), "value", "CareerInfor.name"))
        with self.assertRaisesRegex(ValueError, "不唯一"):
            self.service.prepare(self.request())

    def test_invalid_skill_reference_and_unknown_field_block(self):
        for edits in ({"career": {"unknown": ["1"]}},
                      {"career": {"initial_skill": ["{", "999999;1", "", "}"]}},
                      {"career": {"initial_skill": ["{", "100001;0", "", "}"]}}):
            with self.assertRaises(ValueError):
                self.service.prepare(self.request(edits=edits))

    def test_repeated_slots_are_preserved(self):
        result = self.service.prepare(self.request(edits={"career": {
            "initial_skill": ["{", "100001;1", "100001;1", "}"]}}))
        change = result["plans"][0]["changes"][0]
        self.assertEqual(change["column"], 7)
        self.assertEqual(change["after"], "100001;1")

    def test_clone_copies_row_and_rejects_existing_target(self):
        result = self.service.prepare(self.request(targetId="101"))
        self.assertTrue(result["plans"][0]["clone"])
        self.assertEqual(result["plans"][0]["targetId"], "101")
        self.change_workbook(lambda sheet: sheet.append(["测试", "101"]))
        with self.assertRaisesRegex(ValueError, "已存在"):
            self.service.prepare(self.request(targetId="101"))

    def test_review_is_one_shot_and_bound_to_server_plan(self):
        token = self.service.prepare(self.request())["token"]
        self.service.commit(token)
        self.assertEqual(self.calls[0][1], "copy")
        self.assertEqual(self.calls[0][0][0]["changes"][0]["after"], "测试改名")
        with self.assertRaisesRegex(ValueError, "失效"):
            self.service.commit(token)

    def test_failed_load_preserves_current_repository(self):
        with self.assertRaises(ValueError):
            self.service.load(self.root / "missing")
        self.assertEqual(self.service.detail("100")["name"], "合成职业")

    def test_duplicate_headers_are_retained_positionally(self):
        path = self.root / "duplicate.csv"
        path.write_text("##&X.id,X.buffid,X.buffid\nid,a,b\n1,2,3\n", encoding="utf-8")
        table = Table(path, "X")
        self.assertEqual(table.headers, ["X.id", "X.buffid", "X.buffid"])
        self.assertEqual(table.rows[0], ["1", "2", "3"])
        with self.assertRaises(ValueError):
            table.column("buffid")


if __name__ == "__main__":
    unittest.main()
