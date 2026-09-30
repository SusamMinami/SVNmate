import csv
from pathlib import Path
import tempfile
import unittest

import openpyxl

from backend.authoring import bundle, clone_bundle, validate_operations
from backend.catalog import TABLES
from backend.repository import fingerprint
from backend.service import Service
from test_pipeline import fixture


def authoring_fixture(root):
    fixture(root)
    specs = {
        "skill": (["Skill.id", "Skill.skillname", "Skill.career", "Skill.skillcd", "Skill.subskills"],
                  [["100001", "合成技能", "-1", "1000", ""]]),
        "tree": (["Skillsystem.id", "Skillsystem.skillid", "Skillsystem.skillmaxlv", "Skillsystem.chiefskill"],
                 [["1000001", "100001", "2", "0"]]),
        "upgrade": (["Skillupgrade.id", "Skillupgrade.needsp", "Skillupgrade.para", "", "", ""],
                    [["100001001", "1", "{", "10", "", "}"], ["100001002", "2", "{", "20", "", "}"]]),
        "buff": (["Buffbase.id", "Buffbase.name", "Buffbase.time", "Buff.behavior", "Buff.nextbuff"],
                 [["100001", "同号但独立的Buff", "1000", "", ""]]),
        "behavior": (["Behavior.id", "Behavior.buff", "Behavior.followingBehavior"],
                     [["1", "100001", ""]]),
        "damage": (["Skilldamage.id", "Skilldamage.skillbuff",
                    "SkillBuffInfor.buffid", "SkillBuffInfor.buffhit", "SkillBuffInfor.buffformula", "SkillBuffInfor.bufftime",
                    "SkillBuffInfor.buffid", "SkillBuffInfor.buffhit", "SkillBuffInfor.buffformula", "SkillBuffInfor.bufftime",
                    "SkillBuffInfor.buffid", "SkillBuffInfor.buffhit", "SkillBuffInfor.buffformula", "SkillBuffInfor.bufftime", ""],
                   [["1", "{", "100001", "1", "1", "100", "0", "0", "0", "0", "0", "0", "0", "0", "}"]]),
    }
    for key, (headers, rows) in specs.items():
        stem = TABLES[key][0]
        with (root / "csvdir" / f"{stem}.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows([headers, headers, *rows])
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["版本", *headers, "保留公式"])
        sheet.append(["版本", *headers, "保留公式"])
        for row in rows:
            sheet.append(["合成", *row, "=1+1"])
        book.save(root / "xlsdir" / f"{stem}.xlsx")
        book.close()


class AuthoringTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        authoring_fixture(self.root)
        self.service = Service()
        self.service.load(self.root)
        self.repo = self.service.repo

    def tearDown(self):
        self.tmp.cleanup()

    def cloned(self):
        return clone_bundle(self.repo, {"table": "skill", "sourceId": "100001", "targetId": "100002",
                                       "careerId": "100", "treeIds": {"1000001": "1000002"}})

    def test_skill_clone_covers_tree_and_all_levels_without_same_id_buff(self):
        operations = self.cloned()
        self.assertEqual([o["table"] for o in operations], ["skill", "tree", "upgrade", "upgrade"])
        self.assertEqual([o["targetId"] for o in operations], ["100002", "1000002", "100002001", "100002002"])
        result = self.service.prepare({"sourceId": "100", "operations": operations})
        self.assertEqual(len(result["plans"]), 4)
        self.assertEqual(len({p["path"] for p in result["plans"]}), 3)

    def test_partial_clone_and_bad_level_remap_are_rejected(self):
        operations = self.cloned()
        with self.assertRaisesRegex(ValueError, "全部技能树与升级"):
            validate_operations(self.repo, operations[:-1])
        operations[-1]["targetId"] = "100002003"
        with self.assertRaisesRegex(ValueError, "升级 ID"):
            validate_operations(self.repo, operations)

    def test_new_skill_can_be_attached_to_career_in_same_batch(self):
        result = self.service.prepare({"sourceId": "100", "operations": self.cloned(),
                                       "edits": {"career": {"initial_skill": ["{", "100002;1", "", "}"]}}})
        self.assertEqual(len(result["plans"]), 5)

    def test_buff_full_members_and_repeated_damage_slots_map_to_excel(self):
        buff = bundle(self.repo, "buff", "100001")[0]
        buff["edits"] = {"Buffbase.name": ["合成新名"], "Buff.behavior": ["1"]}
        damage = bundle(self.repo, "damage", "1")[0]
        values = damage["groups"][0]["fields"][0]["values"][:]
        values[5] = "100001"
        damage["edits"] = {"Skilldamage.skillbuff": values}
        result = self.service.prepare({"sourceId": "100", "operations": [buff, damage]})
        self.assertEqual(result["plans"][0]["changes"][1]["member"], "Buff.behavior")
        self.assertEqual(result["plans"][1]["changes"][0]["column"], 8)
        self.assertEqual(values[-1], "}")

    def test_new_buff_reference_resolves_across_draft_and_not_skill_namespace(self):
        buff = bundle(self.repo, "buff", "100001")[0]
        buff["targetId"] = "100002"
        behavior = bundle(self.repo, "behavior", "1")[0]
        behavior["edits"] = {"Behavior.buff": ["100002"]}
        self.assertEqual(len(validate_operations(self.repo, [buff, behavior])), 2)
        with self.assertRaisesRegex(ValueError, "引用 buff 100002"):
            validate_operations(self.repo, [behavior])

    def test_duplicate_targets_unknown_fields_and_array_width_rejected(self):
        operation = bundle(self.repo, "buff", "100001")[0]
        with self.assertRaisesRegex(ValueError, "重复目标"):
            validate_operations(self.repo, [operation, operation])
        operation["edits"] = {"Buffbase.not_real": ["x"]}
        with self.assertRaisesRegex(ValueError, "字段或数组"):
            validate_operations(self.repo, [operation])
        damage = bundle(self.repo, "damage", "1")[0]
        damage["edits"] = {"Skilldamage.skillbuff": ["{", "}"]}
        with self.assertRaisesRegex(ValueError, "字段或数组"):
            validate_operations(self.repo, [damage])

    def test_snapshot_includes_skill_and_buff_files(self):
        previous = self.repo.snapshot
        path = self.root / "csvdir" / "buff表.csv"
        path.write_text(path.read_text("utf-8-sig").replace("1000", "2000"), encoding="utf-8-sig")
        self.service.load(self.root)
        self.assertNotEqual(previous, self.service.repo.snapshot)

    def test_prepare_never_writes_source_files(self):
        before = {p: fingerprint(p) for p in (self.root / "xlsdir").glob("*.xlsx")}
        self.service.prepare({"sourceId": "100", "operations": self.cloned()})
        self.assertEqual(before, {p: fingerprint(p) for p in before})

    def test_legacy_sheet_ignored_but_two_complete_schemas_block(self):
        path = self.root / "xlsdir" / "j技能升级消耗表.xlsx"
        book = openpyxl.load_workbook(path)
        book.create_sheet("历史过渡").append(["Skillupgrade.id", "Skillupgrade.effectperlv"])
        book.save(path)
        book.close()
        self.service.prepare({"sourceId": "100", "operations": self.cloned()})
        book = openpyxl.load_workbook(path)
        book.copy_worksheet(book.worksheets[0])
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, "无法唯一定位"):
            self.service.prepare({"sourceId": "100", "operations": self.cloned()})


if __name__ == "__main__":
    unittest.main()
